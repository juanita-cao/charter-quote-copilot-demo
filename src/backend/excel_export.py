"""E23 · Workbook builder (T2.22, design_backend.md §5.2).

One summary sheet (a row per record) and one detail sheet per record. The detail sheet is an
audit view: its formulas reproduce E2 (`calculate_tce`) step by step, including the lashing
cost and, in display precision, ROUND(...,2) after every step (Excel's ROUND is the same
half-away-from-zero rule as `excel_round`), so what an operator re-adds in Excel equals what
the app showed.

E2 itself is called only to decide which precision a saved record used (the mode that
reproduces its stored TCE); no business arithmetic is re-implemented here beyond the
formula text.
"""

from __future__ import annotations

import io
import json
import re
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.backend.e_nodes import calculate_tce
from src.backend.legacy import MARKER, upgrade_legacy_snapshot
from src.backend.schemas import QuoteInput

_NAVY = "125993"
_LIGHT_BLUE = "D6E4F7"
_LIGHT_GREY = "F2F3F5"
_WHITE = "FFFFFF"
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_TOLERANCE = 0.005  # a stored TCE is kept to 2 decimals


def _f(v: Any, default: float = 0.0) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _snapshot(record: dict) -> dict:
    raw = (
        record.get("raw_input_json")
        if record.get("_kind") == "draft"
        else record.get("quote_input_snapshot")
    )
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except ValueError:
            raw = None
    return upgrade_legacy_snapshot(raw) if isinstance(raw, dict) else {}


def _when(record: dict) -> datetime | None:
    ts = record.get("updated_at") if record.get("_kind") == "draft" else record.get("created_at")
    if isinstance(ts, datetime):
        return ts
    try:
        return datetime.fromisoformat(str(ts))
    except ValueError:
        return None


def _validated(snapshot: dict) -> QuoteInput | None:
    try:
        return QuoteInput.model_validate(snapshot)
    except Exception:
        return None


# ── cell helpers ──────────────────────────────────────────────────────────────


def _style(c, *, bold=False, color="000000", fill=_WHITE, align="left", size=10, fmt=None):
    c.font = Font(bold=bold, color=color, size=size)
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.border = _THIN
    if fmt:
        c.number_format = fmt


def _laycan(snap: dict) -> str:
    start, end = snap.get("laycan_start"), snap.get("laycan_end")
    return f"{str(start)[:10]} – {str(end)[:10]}" if start and end else ""


# ── summary sheet ─────────────────────────────────────────────────────────────

_SUMMARY_COLS = [
    ("Date 保存时间", 20),
    ("Route 航线", 30),
    ("Cargo 货物描述", 16),
    ("Qty · 数量 (RT)", 14),
    ("Rate · 单吨运费", 14),
    ("Total Freight · 总运费 (USD)", 18),
    ("Comm. · 佣金率 (%)", 12),
    ("TCE (USD/day)", 14),
    ("Decision · 决策", 12),
    ("Margin · 利润率 (%)", 14),
    ("Load port · 装港", 16),
    ("Discharge port · 卸港", 16),
    ("Laycan · 受载期", 24),
]


def _write_summary(wb: openpyxl.Workbook, records: list[dict]) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.merge_cells(f"A1:{get_column_letter(len(_SUMMARY_COLS))}1")
    _style(ws["A1"], bold=True, color=_WHITE, fill=_NAVY, align="center", size=13)
    ws["A1"].value = "Voyage quote history · 航次报价历史汇总"
    ws.row_dimensions[1].height = 24
    for i, (label, width) in enumerate(_SUMMARY_COLS, start=1):
        _style(
            ws.cell(row=2, column=i, value=label),
            bold=True,
            color=_WHITE,
            fill=_NAVY,
            align="center",
        )
        ws.column_dimensions[get_column_letter(i)].width = width

    for r, rec in enumerate(records, start=3):
        snap = _snapshot(rec)
        draft = rec.get("_kind") == "draft"
        when = _when(rec)
        qty = snap.get("quantity", rec.get("quantity"))
        rate = snap.get("freight_rate", rec.get("freight_rate"))
        both = isinstance(qty, (int, float)) and isinstance(rate, (int, float))
        row = [
            when.strftime("%Y-%m-%d %H:%M") if when else "",
            rec.get("route") or snap.get("route", ""),
            snap.get("cargo_description", rec.get("cargo_description")) or "",
            _f(qty) if qty is not None else "",
            _f(rate) if rate is not None else "",
            _f(qty) * _f(rate) if both else "—",
            _f(snap.get("commission_rate", rec.get("commission_rate"))),
            "—" if draft else _f(rec.get("tce")),
            "—" if draft else (rec.get("decision") or ""),
            "—" if draft else _f(rec.get("profit_margin_pct")),
            snap.get("load_port") or "",
            snap.get("discharge_port") or "",
            _laycan(snap),
        ]
        fill = _WHITE if r % 2 == 0 else _LIGHT_GREY
        for i, v in enumerate(row, start=1):
            fmt = "#,##0.00" if i in (4, 5, 6, 8) else ("0.00" if i in (7, 10) else None)
            _style(
                ws.cell(row=r, column=i, value=v),
                fill=fill,
                align="right" if isinstance(v, (int, float)) else "left",
                fmt=fmt,
            )


# ── detail sheet ──────────────────────────────────────────────────────────────


class _Detail:
    """Writes label / value / unit rows and remembers the address of each named value."""

    def __init__(self, ws):
        self.ws, self.row, self.ref = ws, 3, {}
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 12
        ws.merge_cells("A1:C1")
        ws["A1"].value = "Voyage estimate · 航次租船预算表"
        _style(ws["A1"], bold=True, color=_WHITE, fill=_NAVY, align="center", size=14)
        ws.row_dimensions[1].height = 28
        for i, label in enumerate(["Item · 项目", "Value · 数值", "Unit · 单位"], start=1):
            _style(ws.cell(row=2, column=i, value=label), bold=True, color=_WHITE, fill=_NAVY)

    def header(self, text: str) -> None:
        for i in range(1, 4):
            _style(
                self.ws.cell(row=self.row, column=i, value=text if i == 1 else None),
                bold=True,
                color=_WHITE,
                fill=_NAVY,
            )
        self.row += 1

    def line(
        self, key: str | None, label: str, value: Any, unit: str = "", fmt: str | None = None
    ) -> str:
        r = self.row
        _style(self.ws.cell(row=r, column=1, value=label), fill=_LIGHT_BLUE)
        _style(self.ws.cell(row=r, column=2, value=value), fmt=fmt)
        _style(self.ws.cell(row=r, column=3, value=unit), color="86909C", size=9)
        addr = f"B{r}"
        if key:
            self.ref[key] = addr
        self.row += 1
        return addr

    def note(self, label: str, text: str) -> None:
        """A long, wrapped text (the cargo notes), merged across the value and unit columns."""
        r = self.row
        _style(self.ws.cell(row=r, column=1, value=label), fill=_LIGHT_BLUE)
        self.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = self.ws.cell(row=r, column=2, value=text)
        _style(c)
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        self.ws.row_dimensions[r].height = 15 * (text.count("\n") + 1 + len(text) // 40)
        self.row += 1

    def blank(self) -> None:
        self.row += 1


_BUCKETS = [  # key, label, laden?
    ("cjk_laden", "CJK Laden · 长江口满载", True),
    ("cjk_ballast", "CJK Ballast · 长江口空载", False),
    ("qz_laden", "QZ Laden · 琼州海峡满载", True),
    ("qz_ballast", "QZ Ballast · 琼州海峡空载", False),
]
_OTHER_COSTS = (
    "loading_cost",
    "discharging_cost",
    "cev_cost",
    "ilohc_cost",
    "lashing_cost",
)
_MONEY, _TWO = "#,##0.00", "0.00"


def _precision_mode(inputs: QuoteInput, stored_tce: float | None) -> str:
    """The mode whose E2 result reproduces the stored TCE (display first), else display."""
    if stored_tce is None:
        return "display"
    if abs(calculate_tce(inputs, "display").tce - stored_tce) <= _TOLERANCE:
        return "display"
    try:
        if abs(calculate_tce(inputs, "full").tce - stored_tce) <= _TOLERANCE:
            return "full"
    except ValueError:
        pass
    return "display"


def _write_calculated(d: _Detail, rec: dict, inputs: QuoteInput, mode: str) -> None:
    def R(expr: str) -> str:  # display precision rounds after every step, like E2
        return f"ROUND({expr},2)" if mode == "display" else expr

    i = inputs.model_dump()
    ref = d.ref
    d.header("How this sheet is calculated · 计算方式")
    d.line(
        "precision",
        "Precision · 计算精度",
        "Display precision — each step rounded to 2 decimals"
        if mode == "display"
        else "Full precision — no rounding",
    )
    if _snapshot(rec).get(MARKER):
        d.line(
            None,
            "PDA split · 港使费拆分",
            "Estimated: the old single PDA was split evenly between the two ports",
        )
    d.blank()

    d.header("Cargo & Freight · 货物与运费")
    d.line("quantity", "Quantity · 货物数量", i["quantity"], "RT")
    d.line("freight_rate", "Freight Rate · 单吨运费", i["freight_rate"], "USD/RT", _MONEY)
    d.line("commission_rate", "Commission Rate · 佣金率", i["commission_rate"], "%", _TWO)
    d.blank()

    d.header("Voyage Time & Distance · 航行时间与距离")
    d.line("bd", "Ballast Distance · 空驶距离", i["ballast_distance"], "nm")
    d.line("ld", "Laden Distance · 满载距离", i["laden_distance"], "nm")
    d.line("bs", "Ballast Speed · 空驶航速", i["ballast_speed"], "kn")
    d.line("ls", "Laden Speed · 满载航速", i["laden_speed"], "kn")
    for key, label, _ in _BUCKETS:
        d.line(f"{key}_nm", f"{label} (nm)", i[f"{key}_nm"], "nm")
        d.line(f"{key}_mgo", f"{label} MGO · 耗轻油", i[f"{key}_mgo_consumption"], "t/d", _TWO)
    for key, en, zh, port_mode, days, rate in (
        ("load", "Loading", "装货", i["loading_mode"], i["loading_days"], i["loading_rate"]),
        (
            "disch",
            "Discharging",
            "卸货",
            i["discharging_mode"],
            i["discharging_days"],
            i["discharging_rate"],
        ),
    ):
        if port_mode == "rate":
            d.line(None, f"{en} Mode · {zh}方式", "Rate: the days are calculated from the rate")
            d.line(f"{key}_rate", f"{en} Rate · {zh}率", rate, "RT/day", _TWO)
            d.line(
                f"{key}_days",
                f"{en} Days · {zh}天数 (from the rate)",
                "=" + R(f"{ref['quantity']}/{ref[f'{key}_rate']}"),
                "days",
                _TWO,
            )
        else:
            d.line(None, f"{en} Mode · {zh}方式", "Days: typed by the operator")
            d.line(f"{key}_days", f"{en} Days · {zh}天数", days, "days", _TWO)
    d.line("margin_days", "Margin Days · 富余天数", i["margin_days"], "days")
    d.blank()

    d.header("Bunker Parameters · 燃油参数")
    for key, label, unit, fmt in [
        ("hfo_price", "HFO Price · 重油价格", "USD/MT", _MONEY),
        ("mgo_price", "MGO Price · 轻油价格", "USD/MT", _MONEY),
        ("hfo_b", "HFO Ballast · 空驶重油耗", "t/d", _TWO),
        ("hfo_l", "HFO Laden · 满载重油耗", "t/d", _TWO),
        ("mgo_b", "MGO Ballast · 空驶轻油耗", "t/d", _TWO),
        ("mgo_l", "MGO Laden · 满载轻油耗", "t/d", _TWO),
        ("hfo_pt", "HFO Port · 在港重油耗", "t/d", _TWO),
        ("mgo_pt", "MGO Port · 在港轻油耗", "t/d", _TWO),
    ]:
        src = {
            "hfo_price": "hfo_price",
            "mgo_price": "mgo_price",
            "hfo_b": "hfo_ballast_consumption",
            "hfo_l": "hfo_laden_consumption",
            "mgo_b": "mgo_ballast_consumption",
            "mgo_l": "mgo_laden_consumption",
            "hfo_pt": "hfo_port_consumption",
            "mgo_pt": "mgo_port_consumption",
        }[key]
        d.line(key, label, i[src], unit, fmt)
    d.blank()

    d.header("Port & Other Cost · 港口与其他成本")
    d.line("load_pda", "Load-port PDA · 装港使费", i["load_port_pda"], "USD", _MONEY)
    d.line("disch_pda", "Discharge-port PDA · 卸港使费", i["discharge_port_pda"], "USD", _MONEY)
    for key, label, src in [
        ("loading_cost", "Loading Cost · 装货费", "loading_cost"),
        ("discharging_cost", "Discharging Cost · 卸货费", "discharging_cost"),
        ("cev_cost", "CEV · 通讯费", "cev_cost"),
        ("ilohc_cost", "ILOHC · 船员扫舱费", "ilohc_cost"),
        ("lashing_cost", "Lashing Cost · 绑扎费", "lashing_cost"),
    ]:
        d.line(key, label, i[src], "USD", _MONEY)
    other_refs = []
    for n, item in enumerate(i["other_costs"], start=1):
        other_refs.append(
            d.line(f"other_{n}", f"Other {n} · {item['name']}", item["amount"], "USD", _MONEY)
        )
    d.blank()

    d.header("Calculation · 计算 (mirrors the app's steps)")
    # Special-passage days first: in display precision a leg's days are its own rounded figure
    # minus these rounded days (ADR-023 amendment 2026-09-24, the client's workbook).
    special_days, special_mgo = [], []
    for key, label, laden in _BUCKETS:
        speed = ref["ls"] if laden else ref["bs"]
        days = d.line(
            f"{key}_days",
            f"{label} Days · 天数",
            "=" + R(f"{ref[f'{key}_nm']}/{speed}/24"),
            "days",
            _TWO,
        )
        qty = d.line(
            f"{key}_mgo_qty",
            f"{label} MGO Qty · 耗轻油量",
            "=" + R(f"{ref[f'{key}_mgo']}*{days}"),
            "t",
            _TWO,
        )
        special_days.append(days)
        special_mgo.append(qty)

    def leg_days(dist: str, speed: str, nm_keys: tuple[str, str], day_cells: list[str]) -> str:
        if mode != "display":
            special = f"SUM({ref[nm_keys[0]]},{ref[nm_keys[1]]})"
            return f"=({ref[dist]}-{special})/{ref[speed]}/24"
        return f"=MAX(0,ROUND(ROUND({ref[dist]}/{ref[speed]}/24,2)-SUM({','.join(day_cells)}),2))"

    # _BUCKETS order: cjk_laden, cjk_ballast, qz_laden, qz_ballast
    d.line(
        "bdays",
        "Ballast Days · 空驶天数",
        leg_days(
            "bd", "bs", ("cjk_ballast_nm", "qz_ballast_nm"), [special_days[1], special_days[3]]
        ),
        "days",
        _TWO,
    )
    d.line(
        "ldays",
        "Laden Days · 满载天数",
        leg_days("ld", "ls", ("cjk_laden_nm", "qz_laden_nm"), [special_days[0], special_days[2]]),
        "days",
        _TWO,
    )
    d.line(
        "port_days",
        "Port Days · 在港天数",
        "=" + R(f"{ref['load_days']}+{ref['disch_days']}+{ref['margin_days']}"),
        "days",
        _TWO,
    )
    d.line(
        "total_days",
        "Total Voyage Days · 航次总天数",
        "=" + R(f"{ref['bdays']}+{ref['ldays']}+SUM({','.join(special_days)})+{ref['port_days']}"),
        "days",
        _TWO,
    )

    def leg(consumption: str, days: str) -> str:
        return R(f"{ref[consumption]}*{ref[days]}")

    d.line(
        "hfo_qty",
        "Total HFO · 重油总耗",
        "=" + "+".join([leg("hfo_b", "bdays"), leg("hfo_l", "ldays"), leg("hfo_pt", "port_days")]),
        "t",
        _TWO,
    )
    d.line(
        "mgo_qty",
        "Total MGO · 轻油总耗",
        "="
        + "+".join(
            [
                leg("mgo_b", "bdays"),
                leg("mgo_l", "ldays"),
                leg("mgo_pt", "port_days"),
                f"SUM({','.join(special_mgo)})",
            ]
        ),
        "t",
        _TWO,
    )
    d.line(
        "hfo_cost",
        "HFO Cost · 重油成本",
        "=" + R(f"{ref['hfo_qty']}*{ref['hfo_price']}"),
        "USD",
        _MONEY,
    )
    d.line(
        "mgo_cost",
        "MGO Cost · 轻油成本",
        "=" + R(f"{ref['mgo_qty']}*{ref['mgo_price']}"),
        "USD",
        _MONEY,
    )
    d.line(
        "bunker",
        "Total Bunker · 总燃油成本",
        "=" + R(f"{ref['hfo_cost']}+{ref['mgo_cost']}"),
        "USD",
        _MONEY,
    )
    d.line(
        "pda_total",
        "PDA Total · 港使费合计",
        "=" + R(f"{ref['load_pda']}+{ref['disch_pda']}"),
        "USD",
        _MONEY,
    )
    others_formula = "=" + R(f"SUM({','.join(other_refs)})") if other_refs else 0
    d.line("others_total", "Other Costs Total · 其他费用合计", others_formula, "USD", _MONEY)
    d.line(
        "voyage_cost",
        "Total Voyage Cost · 航次总成本",
        "="
        + R(
            f"{ref['bunker']}+{ref['pda_total']}"
            f"+SUM({','.join(ref[k] for k in _OTHER_COSTS)})+{ref['others_total']}"
        ),
        "USD",
        _MONEY,
    )
    d.line(
        "revenue",
        "Total Freight · 总运费",
        "=" + R(f"{ref['quantity']}*{ref['freight_rate']}"),
        "USD",
        _MONEY,
    )
    d.line(
        "commission",
        "Commission · 佣金",
        "=" + R(f"{ref['revenue']}*{ref['commission_rate']}/100"),
        "USD",
        _MONEY,
    )
    d.line(
        "net_freight",
        "Net Freight · 净运费收入",
        "=" + R(f"{ref['revenue']}-{ref['commission']}"),
        "USD",
        _MONEY,
    )
    d.line(
        "net_income",
        "Net Voyage Income · 航次净收入",
        "=" + R(f"{ref['net_freight']}-{ref['voyage_cost']}"),
        "USD",
        _MONEY,
    )
    d.line(
        "tce",
        "TCE (USD/day)",
        "=" + R(f"{ref['net_income']}/{ref['total_days']}"),
        "USD/day",
        _MONEY,
    )
    d.blank()


def _write_stored(d: _Detail, rec: dict, has_formulas: bool) -> None:
    if rec.get("_kind") == "draft":
        return
    d.header("Saved with the quote · 保存时的结果")
    stored = d.line("stored_tce", "Stored TCE · 保存的 TCE", _f(rec.get("tce")), "USD/day", _MONEY)
    if has_formulas:
        d.line(
            "diff",
            "Difference · 差异 (recalculated − stored)",
            f"={d.ref['tce']}-{stored}",
            "USD/day",
            _MONEY,
        )
    d.line("decision", "Stored Decision · 保存的决策", rec.get("decision") or "")
    d.line(
        "margin", "Stored Profit Margin · 保存的利润率", _f(rec.get("profit_margin_pct")), "%", _TWO
    )
    d.line(
        "ask",
        "Shipowner Asking TCE · 船东要价",
        _f(_snapshot(rec).get("shipowner_asking_tce")),
        "USD/day",
        _MONEY,
    )
    d.line(
        "spread",
        "Hire Spread vs Owner Ask · 与船东租金价差",
        f"={stored}-{d.ref['ask']}",
        "USD/day",
        _MONEY,
    )


def _write_inputs_only(d: _Detail, rec: dict) -> None:
    d.header("Draft — inputs as entered, not calculated · 草稿：仅录入内容，未计算")
    for key, value in _snapshot(rec).items():
        if isinstance(value, (str, int, float, bool)):
            d.line(None, str(key), value)


def _write_detail(wb: openpyxl.Workbook, rec: dict, name: str) -> None:
    ws = wb.create_sheet(title=name)
    d = _Detail(ws)
    d.header("Basic Info · 基本信息")
    when = _when(rec)
    d.line(None, "Saved At · 保存时间", when.strftime("%Y-%m-%d %H:%M") if when else "")
    d.line(None, "Route · 航线", rec.get("route") or _snapshot(rec).get("route", ""))
    d.line(
        None,
        "Cargo · 货物描述",
        rec.get("cargo_description") or _snapshot(rec).get("cargo_description", ""),
    )
    snap = _snapshot(rec)
    d.line(None, "Fill-in Date · 填表日期", str(snap.get("fill_in_date") or "")[:10])
    d.line(None, "Laycan · 受载期", _laycan(snap))
    d.line(None, "Load Port · 装港", snap.get("load_port") or "")
    d.line(None, "Discharge Port · 卸港", snap.get("discharge_port") or "")
    if snap.get("cargo_notes"):
        d.note("Cargo Notes · 货盘备注", snap["cargo_notes"])
    d.blank()

    inputs = _validated(_snapshot(rec))
    stored = (
        _f(rec.get("tce")) if rec.get("_kind") != "draft" and rec.get("tce") is not None else None
    )
    if inputs is not None:
        try:
            _write_calculated(d, rec, inputs, _precision_mode(inputs, stored))
            _write_stored(d, rec, True)
        except ValueError:  # e.g. a special passage larger than its leg: E2 refuses it too
            _write_stored(d, rec, False)
    else:
        if rec.get("_kind") == "draft":
            _write_inputs_only(d, rec)
        else:
            _write_stored(d, rec, False)
    ws.freeze_panes = "A3"


_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(rec: dict, taken: set[str]) -> str:
    when = _when(rec)
    stamp = when.strftime("%m-%d") if when else "??-??"
    prefix = "Draft " if rec.get("_kind") == "draft" else ""
    base = _BAD_SHEET_CHARS.sub(
        "-",
        f"{prefix}{stamp} {(rec.get('route') or 'quote')[:12]}",
    ).strip()[:31]
    name, n = base, 1
    while name in taken:
        n += 1
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
    taken.add(name)
    return name


def build_history_workbook(records: list[dict]) -> bytes:
    """E23: the `.xlsx` bytes, records listed newest first."""
    ordered = sorted(records, key=lambda r: _when(r) or datetime.min, reverse=True)
    wb = openpyxl.Workbook()
    _write_summary(wb, ordered)
    taken = {"Summary"}
    for rec in ordered:
        _write_detail(wb, rec, _sheet_name(rec, taken))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
